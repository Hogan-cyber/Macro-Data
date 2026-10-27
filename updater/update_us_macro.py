#!/usr/bin/env python3
"""Refresh the US macro workbook. Python 3.10+, openpyxl, certifi.

No API key is required: uses FRED's official public CSV download and series
metadata pages. This is a latest-revised dataset, not a historical PIT dataset.
"""
from __future__ import annotations

import argparse
import concurrent.futures
import csv
import hashlib
import html
import io
import json
import math
import os
from pathlib import Path
import re
import shutil
import ssl
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parent
UTC = timezone.utc
DATA_HEADERS = ['observation_date', 'value', 'is_missing']
DICT_HEADERS = ['series_id','sheet_name','indicator_zh','category','frequency','unit',
                'seasonal_adjustment','source_title','source_url','download_url',
                'first_observation','last_observation','valid_rows','missing_rows',
                'retrieved_at_utc','source_last_updated','update_status','definition',
                'usage_note','method','input_series_id','error']
LOG_HEADERS = ['run_id','completed_at_utc','status','source_series','sources_updated',
               'sources_failed','new_observations','revised_observations','total_valid_rows']
REV_HEADERS = ['run_id','detected_at_utc','series_id','observation_date','old_value','new_value','change_type']
NOTE_ROWS = [
    ['用途','美国宏观与市场数据，按指标分表。数据区固定为 A:C，从第一行表头开始。'],
    ['日期','observation_date 是观测期间标签，不是发布日期。月度/季度通常标记该期首日。'],
    ['数值','value 为原始来源单位的数值；利率 4.5 表示 4.5%，不是 0.045。单位见数据页右侧和数据字典。'],
    ['缺失','空值保持为空，is_missing=1；真实零值仍为 0，is_missing=0。不做插值或向前/向后填充。'],
    ['历史','默认从 2000-01-01 开始。各序列历史长度按真实可用情况保留；部分市场数据仅有较短历史。'],
    ['修订','当前值是最新修订历史数据。以后每次运行会比较已保存值，并在修订记录保留旧值和新值。'],
    ['回测限制','本文件不包含完整历史首次发布值或可靠逐期发布时间，不能直接宣称为无前视偏差的训练集。'],
    ['派生','同比按同一日历月份的上年值计算；非农月度新增为相邻月份总人数之差。缺少对应月份则留空。'],
    ['就业口径','就业人口比率的分母为成年非机构平民人口；失业率的分母为劳动力，两者不是互补关系。'],
    ['数据来源','使用 FRED 官方公开下载与元数据页面；每项原始来源标题、单位、频率和链接可追溯。'],
    ['自动更新','运行 update_us_macro.py 即抓取并更新；--watch 会持续等待香港时间每周一 08:00，需机器和程序保持运行。'],
    ['后台计划','仓库 .github/workflows/weekly.yml 配置云端更新；请到 Actions 查看运行状态。--watch 仅用于本地。'],
    ['备份','更新前保留原工作簿至 backups/，新文件校验成功后原子替换。原始响应保存在 data/raw/。'],
    ['失败处理','单项失败保留上一合格数据，并标记 FAILED_RETAINED。首次获取失败标记 FAILED_EMPTY，不填假值。'],
    ['编辑','更新程序维护已有指标页和说明/目录/日志页。新增的个人工作表会保留；请把手工模型放在独立工作簿。'],
    ['读取','pandas.read_excel("美国宏观数据.xlsx", sheet_name="非农月度新增", usecols="A:C", parse_dates=["observation_date"])'],
    ['原始频率','日、周、月、季度分别保留。每周运行更新程序不会把所有指标改成周频。'],
    ['异常值','真实极端值不自动删除；例如原油负价格可以是有效观察。'],
    ['来源权限','部分指数和利差的原始提供者保留版权。公开下载不等于可公开再分发；使用范围见每项来源页面。'],
]

def now_iso():
    return datetime.now(UTC).isoformat(timespec='seconds').replace('+00:00','Z')

def parse_csv(text, series_id):
    reader = csv.DictReader(io.StringIO(text.lstrip('\ufeff')))
    if reader.fieldnames != ['observation_date', series_id]:
        raise ValueError(f'Unexpected CSV header for {series_id}: {reader.fieldnames}')
    values = {}
    for row in reader:
        day = date.fromisoformat(row['observation_date']).isoformat()
        raw = (row[series_id] or '').strip()
        value = None if raw in ('', '.', 'NA', 'N/A') else float(raw)
        if value is not None and not math.isfinite(value):
            raise ValueError(f'Non-finite value at {day}')
        if day in values and values[day] != value:
            raise ValueError(f'Conflicting duplicate at {day}')
        values[day] = value
    if not values or not any(v is not None for v in values.values()):
        raise ValueError(f'Empty or all-missing response for {series_id}')
    return [[d,v,int(v is None)] for d,v in sorted(values.items())]

def derive_rows(rows, method):
    lookup = {r[0]:r[1] for r in rows}
    result = []
    for day, value, _ in rows:
        d = date.fromisoformat(day)
        if d.day != 1:
            raise ValueError('Monthly derivation requires month-start labels')
        prior = date(d.year-1,d.month,1) if method=='yoy' else date(d.year-(d.month==1),12 if d.month==1 else d.month-1,1)
        previous = lookup.get(prior.isoformat())
        calculated = None
        if value is not None and previous is not None:
            if method=='diff':calculated = value-previous
            elif method=='yoy' and previous!=0:calculated = (value/previous-1)*100
            elif method!='yoy':raise ValueError(f'Unknown method {method}')
        result.append([day,calculated,int(calculated is None)])
    return result

def compare_rows(old, new):
    before = {r[0]:r[1] for r in old}
    after = {r[0]:r[1] for r in new}
    removed = before.keys()-after.keys()
    if removed:raise ValueError(f'Source response lost {len(removed)} saved dates; retained previous series')
    revisions=[]
    for day in sorted(before.keys() & after.keys()):
        a,b=before[day],after[day]
        equal = a==b or (a is not None and b is not None and math.isclose(a,b,rel_tol=0,abs_tol=1e-12))
        if not equal:
            kind = 'filled_missing' if a is None else 'now_missing' if b is None else 'revised'
            revisions.append([day,a,b,kind])
    return {'new_rows':len(after.keys()-before.keys()),'changed_rows':len(revisions)},revisions

def retain_rolling_history(old,new):
    """Preserve an archived prefix only for explicitly configured rolling sources.

    Missing dates inside the returned range still fail compare_rows. A response
    with no overlap cannot establish continuity and is rejected.
    """
    if not old:return new
    if not new or not ({r[0] for r in old}&{r[0] for r in new}):
        raise ValueError('Rolling source has no overlap with saved history')
    return [r for r in old if r[0]<new[0][0]]+new

def get_url(url):
    try:
        import certifi
        context=ssl.create_default_context(cafile=certifi.where())
    except ImportError:
        context=ssl.create_default_context()
    for attempt in range(3):
        try:
            req=urllib.request.Request(url)
            with urllib.request.urlopen(req,timeout=45,context=context) as response:
                return response.read().decode('utf-8-sig')
        except (urllib.error.URLError,TimeoutError,OSError):
            if attempt==2:raise
            time.sleep(2**attempt)

def clean_html(text):
    return re.sub(r'\s+',' ',html.unescape(re.sub('<[^>]+>',' ',text))).strip()

def metadata_from_html(page):
    def span(css):
        match=re.search(r'<span[^>]*class="[^"]*\b'+re.escape(css)+r'\b[^"]*"[^>]*>(.*?)</span>',page,re.S)
        return clean_html(match.group(1)) if match else ''
    unit=span('series-meta-value-units')
    frequency=span('series-meta-value-frequency')
    pos=page.find('series-meta-value-units')
    section=clean_html(page[pos:pos+900])
    match=re.search(r'Not Seasonally Adjusted|Seasonally Adjusted Annual Rate|Seasonally Adjusted',section)
    seasonal=match.group(0) if match else ''
    title=re.search(r'<title>(.*?)</title>',page,re.S)
    if not unit or not frequency or not seasonal:
        raise ValueError('Could not verify series units/frequency/seasonal adjustment from source page')
    return {'unit':unit,'frequency':frequency,'seasonal_adjustment':seasonal,
            'source_title':clean_html(title.group(1)).split('| FRED')[0].strip() if title else '',
            'source_last_updated':(span('series-meta-updated-date')+' '+span('series-meta-updated-time')).strip(),
            'usage_note':'Third-party copyright: review source terms' if 'Copyrighted: Pre-Approval Required' in page or re.search(r'Copyright[^<]{0,50}(Nasdaq|NASDAQ|ICE|S&P|Cboe|CBOE)',page) else 'Review original source terms'}

def read_existing(path, cfg):
    if not path.exists():return {},[],[]
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=True)
    metas={}
    if '数据字典' in wb:
        iterator=wb['数据字典'].iter_rows(values_only=True)
        headers=next(iterator)
        for vals in iterator:
            row=dict(zip(headers,vals))
            for key in ('first_observation','last_observation'):
                if isinstance(row.get(key),datetime):row[key]=row[key].date().isoformat()
                elif isinstance(row.get(key),date):row[key]=row[key].isoformat()
            if isinstance(row.get('retrieved_at_utc'),datetime):
                row['retrieved_at_utc']=row['retrieved_at_utc'].replace(tzinfo=UTC).isoformat(timespec='seconds').replace('+00:00','Z')
            if row.get('series_id'):metas[row['series_id']]=row
    previous={}
    for spec in cfg['series']+cfg['derived']:
        if spec['sheet'] not in wb:continue
        ws=wb[spec['sheet']]
        if list(next(ws.iter_rows(min_row=1,max_row=1,max_col=3,values_only=True)))!=DATA_HEADERS:
            raise ValueError(f'Workbook schema mismatch: {spec["sheet"]}')
        rows=[]
        for d,v,missing in ws.iter_rows(min_row=2,max_col=3,values_only=True):
            if d is None:continue
            day=d.date().isoformat() if isinstance(d,datetime) else d.isoformat() if isinstance(d,date) else str(d)
            date.fromisoformat(day)
            if v is not None and (isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v)):
                raise ValueError(f'Invalid saved numeric value: {spec["sheet"]} {day}')
            rows.append([day,v,int(v is None)])
        previous[spec['id']]={'rows':rows,'meta':metas.get(spec['id'],{})}
    logs=list(wb['更新记录'].iter_rows(min_row=2,values_only=True)) if '更新记录' in wb else []
    revisions=list(wb['修订记录'].iter_rows(min_row=2,values_only=True)) if '修订记录' in wb else []
    wb.close()
    def normalize_record(row,is_revision=False):
        result=list(row)
        for i,value in enumerate(result):
            if isinstance(value,datetime):
                result[i]=value.date().isoformat() if is_revision and i==3 else value.replace(tzinfo=UTC).isoformat(timespec='seconds').replace('+00:00','Z')
            elif isinstance(value,date):result[i]=value.isoformat()
        return result
    return previous,[normalize_record(r) for r in logs if r[0]],[normalize_record(r,True) for r in revisions if r[0]]

def fetch_one(spec, fetch_start, end, raw_dir):
    sid=spec['id']
    url='https://fred.stlouisfed.org/graph/fredgraph.csv?'+urllib.parse.urlencode({'id':sid,'cosd':fetch_start,'coed':end})
    data=get_url(url)
    rows=parse_csv(data,sid)
    if any(r[0]>end for r in rows):raise ValueError('Response contains observations after requested cutoff')
    page=get_url('https://fred.stlouisfed.org/series/'+sid)
    meta=metadata_from_html(page)
    (raw_dir/(sid+'.csv')).write_text(data,encoding='utf-8')
    meta.update({'source_url':'https://fred.stlouisfed.org/series/'+sid,'download_url':url,
                 'retrieved_at_utc':now_iso(),'sha256':hashlib.sha256(data.encode()).hexdigest()})
    (raw_dir/(sid+'.metadata.json')).write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
    return rows,meta

def dataset_entry(spec,rows,meta,status,error=''):
    valid=[r for r in rows if r[1] is not None]
    meta={**meta,'series_id':spec['id'],'sheet_name':spec['sheet'],'indicator_zh':spec['label'],
          'category':spec['group'],'first_observation':valid[0][0] if valid else '',
          'last_observation':valid[-1][0] if valid else '', 'valid_rows':len(valid),
          'missing_rows':len(rows)-len(valid),'update_status':status,'error':error,
          'method':spec.get('method','source'),'input_series_id':spec.get('input',''),
          'definition':meta.get('definition') or spec['label']}
    return {'spec':spec,'rows':rows,'meta':meta}

def collect(cfg, output, start, raw_root):
    started=now_iso(); run_id=datetime.now(UTC).strftime('%Y%m%dT%H%M%S%fZ')
    prior,logs,revisions=read_existing(output,cfg)
    # Changing start dates must not silently remove rows from a maintained file.
    for value in prior.values():
        if value['rows'] and value['rows'][0][0]<start:
            raise ValueError('Requested start would remove existing history; use a different output file')
    fetch_start=date(date.fromisoformat(start).year-1,1,1).isoformat()
    end=datetime.now(ZoneInfo('Asia/Hong_Kong')).date().isoformat()
    raw_dir=raw_root/run_id;raw_dir.mkdir(parents=True,exist_ok=False)
    result={}; full_rows={};new_count=0;changed_count=0;failures=[]
    with concurrent.futures.ThreadPoolExecutor(max_workers=3) as pool:
        futures={pool.submit(fetch_one,spec,fetch_start,end,raw_dir):spec for spec in cfg['series']}
        for future in concurrent.futures.as_completed(futures):
            spec=futures[future];sid=spec['id'];old=prior.get(sid,{'rows':[],'meta':{}})
            try:
                full,meta=future.result();rows=[r for r in full if r[0]>=start]
                if not any(r[1] is not None for r in rows):raise ValueError('No numeric history within requested range')
                if spec.get('rolling_archive'):
                    rows=retain_rolling_history(old['rows'],rows)
                    meta['usage_note']+='; earlier saved observations outside the rolling download window are retained'
                for field in ('unit','frequency','seasonal_adjustment'):
                    if old['meta'].get(field) and old['meta'][field]!=meta[field]:
                        raise ValueError(f'Source {field} changed; previous series retained for review')
                counts,changes=compare_rows(old['rows'],rows)
                new_count+=counts['new_rows'];changed_count+=counts['changed_rows']
                revisions.extend([[run_id,started,sid,*change] for change in changes])
                full_rows[sid]=full
                result[sid]=dataset_entry(spec,rows,meta,'UPDATED')
                print(f'{sid}: {len(rows)} rows',flush=True)
            except Exception as exc:
                error=str(exc);failures.append({'series_id':sid,'error':error})
                result[sid]=dataset_entry(spec,old['rows'],old['meta'],'FAILED_RETAINED' if old['rows'] else 'FAILED_EMPTY',error)
                print(f'{sid}: FAILED ({error})',flush=True)
    for spec in cfg['derived']:
        source=result[spec['input']];old=prior.get(spec['id'],{'rows':[],'meta':{}})
        if source['meta']['update_status']!='UPDATED':
            result[spec['id']]=dataset_entry(spec,old['rows'],old['meta'],'FAILED_RETAINED' if old['rows'] else 'FAILED_EMPTY','Input source failed')
            continue
        meta=dict(source['meta'])
        meta['unit']='Percent' if spec['method']=='yoy' else source['meta']['unit']
        meta['source_title']='Derived from '+source['meta']['source_title']
        meta['definition']='100 * (value[t] / value[same month last year] - 1)' if spec['method']=='yoy' else 'value[t] - value[previous calendar month]'
        rows=[r for r in derive_rows(full_rows[spec['input']],spec['method']) if r[0]>=start]
        counts,changes=compare_rows(old['rows'],rows)
        new_count+=counts['new_rows'];changed_count+=counts['changed_rows']
        revisions.extend([[run_id,started,spec['id'],*change] for change in changes])
        result[spec['id']]=dataset_entry(spec,rows,meta,'UPDATED')
    entries=[result[s['id']] for s in cfg['series']+cfg['derived']]
    done=now_iso();status='PARTIAL' if failures else 'SUCCESS'
    logs.append([run_id,done,status,len(cfg['series']),len(cfg['series'])-len(failures),len(failures),new_count,changed_count,sum(e['meta']['valid_rows'] for e in entries)])
    return {'schema_version':1,'run_id':run_id,'started_at_utc':started,'completed_at_utc':done,
            'start_date':start,'status':status,'entries':entries,'logs':logs,'revisions':revisions,
            'failures':failures,'dictionary_headers':DICT_HEADERS,'log_headers':LOG_HEADERS,
            'revision_headers':REV_HEADERS,'notes':NOTE_ROWS}

def write_workbook(snapshot, output):
    # The portable updater deliberately uses Python/openpyxl, as requested.
    from openpyxl import Workbook,load_workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.worksheet.table import Table,TableStyleInfo
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter
    wb=load_workbook(output) if output.exists() else Workbook()
    names=['指标目录']+[e['spec']['sheet'] for e in snapshot['entries']]+['数据字典','更新记录','修订记录','使用说明']
    for name in names:
        if name in wb:wb.remove(wb[name])
    if 'Sheet' in wb and wb['Sheet'].max_row==1 and wb['Sheet']['A1'].value is None:wb.remove(wb['Sheet'])
    navy='17324D';teal='176B74';muted='64748B'
    def style(ws,last_row,last_col,widths,header_row=1):
        ws.sheet_view.showGridLines=False;ws.freeze_panes=f'A{header_row+1}'
        for cells in ws.iter_rows(min_row=1,max_row=last_row,max_col=last_col):
            for c in cells:c.font=Font(name='Arial',size=10);c.alignment=Alignment(vertical='center')
        for c in ws[header_row][:last_col]:
            c.fill=PatternFill('solid',fgColor=navy);c.font=Font(name='Arial',size=10,color='FFFFFF',bold=True)
            c.alignment=Alignment(vertical='center',horizontal='center')
        ws.row_dimensions[header_row].height=26
        for idx,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(idx)].width=width
        if last_row>header_row:ws.auto_filter.ref=f'A{header_row}:{get_column_letter(last_col)}{last_row}'
    index=wb.create_sheet('指标目录',0)
    index['A2']='美国宏观与市场数据';index['A2'].font=Font(name='Arial',size=16,bold=True,color=navy)
    index['A3']='最新修订口径 · '+snapshot['completed_at_utc']+' · 数据区 A:C'
    index['A4']='GitHub Actions 计划：每周一 08:00 香港时间；实际执行状态见 GitHub Actions'
    headers=['工作表','类别','频率','最新观测期','最新值','单位','有效观测数','更新状态']
    for c,v in enumerate(headers,1):index.cell(6,c,v)
    for i,e in enumerate(snapshot['entries'],7):
        m=e['meta'];valid=[r for r in e['rows'] if r[1] is not None];last=valid[-1] if valid else [None,None]
        vals=[e['spec']['sheet'],m['category'],m.get('frequency',''),date.fromisoformat(last[0]) if last[0] else None,last[1],m.get('unit',''),m['valid_rows'],m['update_status']]
        for c,v in enumerate(vals,1):index.cell(i,c,v)
        index.cell(i,1).hyperlink=f"#'{e['spec']['sheet']}'!A1"
        index.cell(i,4).number_format='yyyy-mm-dd';index.cell(i,5).number_format='#,##0.0000'
    style(index,len(snapshot['entries'])+6,8,[25,16,24,18,18,34,17,24],6)
    index['A2'].font=Font(name='Arial',size=16,bold=True,color=navy);index.row_dimensions[2].height=28
    index['A3'].font=Font(name='Arial',size=10,color=muted)
    index['A4'].font=Font(name='Arial',size=10,color=muted)
    index.sheet_properties.tabColor=navy
    for i,e in enumerate(snapshot['entries'],1):
        ws=wb.create_sheet(e['spec']['sheet'],i);m=e['meta'];ws.append(DATA_HEADERS)
        for d,v,flag in e['rows']:ws.append([date.fromisoformat(d),v,flag])
        count=len(e['rows'])+1
        style(ws,count,3,[22,22,17])
        for row in ws.iter_rows(min_row=2,max_row=count,max_col=3):
            row[0].number_format='yyyy-mm-dd';row[1].number_format='#,##0.0000'
        if count>1:
            table=Table(displayName='Data_'+e['spec']['id'],ref=f'A1:C{count}')
            table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
            ws.add_table(table)
            ws.conditional_formatting.add(f'C2:C{count}',CellIsRule(operator='equal',formula=['1'],fill=PatternFill('solid',fgColor='FFF0CF')))
        pairs=[['指标',m['indicator_zh']],['标识',m['series_id']],['单位',m.get('unit','')],['频率',m.get('frequency','')],['季节调整',m.get('seasonal_adjustment','')],['来源',m.get('source_url','')],['获取时间 UTC',m.get('retrieved_at_utc','')],['源站更新时间',m.get('source_last_updated','')],['更新状态',m['update_status']],['数据口径','最新修订值；非发布日期对齐数据'],['定义',m.get('definition','')]]
        for r,pair in enumerate(pairs,1):
            for c,v in enumerate(pair,5):
                ws.cell(r,c,v);ws.cell(r,c).font=Font(name='Arial',size=10,color=muted if c==5 else navy)
        ws.column_dimensions['D'].width=3;ws.column_dimensions['E'].width=22;ws.column_dimensions['F'].width=85
    def flat_sheet(name,headers,rows,widths):
        ws=wb.create_sheet(name);ws.append(headers)
        for row in rows:ws.append(row)
        style(ws,len(rows)+1,len(headers),widths)
        return ws
    flat_sheet('数据字典',DICT_HEADERS,[[e['meta'].get(h,'') for h in DICT_HEADERS] for e in snapshot['entries']],[23,25,48,18,24,38,32,70,65,85,22,22,17,17,28,32,24,75,60,16,24,70])
    flat_sheet('更新记录',LOG_HEADERS,snapshot['logs'],[32,28,18,18,20,20,24,24,24])
    flat_sheet('修订记录',REV_HEADERS,snapshot['revisions'],[32,28,24,22,22,22,22])
    ws=flat_sheet('使用说明',['项目','说明'],NOTE_ROWS,[22,125]);ws.sheet_properties.tabColor='94A3B8'
    for row in ws.iter_rows(min_row=2):
        row[1].alignment=Alignment(vertical='center',wrap_text=True);ws.row_dimensions[row[0].row].height=34
    wb.active=0
    output.parent.mkdir(parents=True,exist_ok=True)
    temp=output.with_name('.'+output.stem+'.pending.xlsx')
    wb.save(temp);wb.close()
    check=load_workbook(temp,read_only=True,data_only=True)
    for entry in snapshot['entries']:
        ws=check[entry['spec']['sheet']]
        actual=sum(1 for row in ws.iter_rows(min_row=2,max_col=1,values_only=True) if row[0] is not None)
        if actual!=len(entry['rows']):raise ValueError('Export row count mismatch')
    check.close()
    if output.exists():
        backup=output.parent/'backups';backup.mkdir(exist_ok=True)
        shutil.copy2(output,backup/(output.stem+'_'+snapshot['run_id']+'.xlsx'))
    os.replace(temp,output)

def acquire_lock(path):
    import contextlib
    @contextlib.contextmanager
    def locked():
        path.parent.mkdir(parents=True,exist_ok=True)
        stream=path.open('a+b')
        try:
            if os.name=='nt':
                import msvcrt
                stream.seek(0);stream.write(b'0');stream.flush();stream.seek(0)
                msvcrt.locking(stream.fileno(),msvcrt.LK_NBLCK,1)
            else:
                import fcntl
                fcntl.flock(stream.fileno(),fcntl.LOCK_EX|fcntl.LOCK_NB)
            yield
        finally:stream.close()
    return locked()

def run(args):
    output=Path(args.output).resolve()
    cfg=json.loads(Path(args.config).read_text(encoding='utf-8'))
    start=args.start or cfg['default_start'];date.fromisoformat(start)
    with acquire_lock(output.with_suffix('.lock')):
        if args.from_snapshot:
            snapshot=json.loads(Path(args.from_snapshot).read_text(encoding='utf-8'))
        else:
            snapshot=collect(cfg,output,start,output.parent/'raw')
        snapshot_path=Path(args.snapshot) if args.snapshot else output.parent/'latest_snapshot.json'
        snapshot_path.parent.mkdir(parents=True,exist_ok=True)
        snapshot_path.write_text(json.dumps(snapshot,ensure_ascii=False,allow_nan=False),encoding='utf-8')
        if not args.collect_only:write_workbook(snapshot,output)
        print(json.dumps({'run_id':snapshot['run_id'],'status':snapshot['status'],'source_failures':len(snapshot['failures']),
                          'sheets':len(snapshot['entries']),'valid_rows':sum(e['meta']['valid_rows'] for e in snapshot['entries']),
                          'workbook_written':not args.collect_only,'output':str(output)},ensure_ascii=False),flush=True)
        return 2 if snapshot['failures'] else 0

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--output',default=str(ROOT.parent/'data/us_macro.xlsx'))
    p.add_argument('--config',default=str(ROOT/'series.json'))
    p.add_argument('--start',help='YYYY-MM-DD; default 2000-01-01')
    p.add_argument('--collect-only',action='store_true',help='Save a JSON snapshot without writing the workbook')
    p.add_argument('--snapshot',help='Optional JSON snapshot destination')
    p.add_argument('--from-snapshot',help='Rebuild from a saved snapshot without network access')
    p.add_argument('--watch',action='store_true',help='Keep running and refresh Mondays 08:00 Asia/Hong_Kong')
    args=p.parse_args()
    if args.watch and (args.collect_only or args.from_snapshot):p.error('--watch cannot be combined with snapshot-only modes')
    code=run(args)
    if not args.watch:return code
    tz=ZoneInfo('Asia/Hong_Kong')
    while True:
        now=datetime.now(tz);target=now.replace(hour=8,minute=0,second=0,microsecond=0)+timedelta(days=(0-now.weekday())%7)
        if target<=now:target+=timedelta(days=7)
        print('Next update: '+target.isoformat(),flush=True)
        while datetime.now(tz)<target:time.sleep(min(60,max(1,(target-datetime.now(tz)).total_seconds())))
        try:run(args)
        except Exception as exc:print('Update failed: '+str(exc),file=sys.stderr,flush=True)

if __name__=='__main__':
    try:sys.exit(main())
    except KeyboardInterrupt:sys.exit(130)
    except Exception as exc:
        print('Update failed: '+str(exc),file=sys.stderr);sys.exit(1)
